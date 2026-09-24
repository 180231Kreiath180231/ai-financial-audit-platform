from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document as WordDocument
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader

from backend.app.db import Database
from backend.app.outputs import (
    OutputDraftService,
    OutputError,
    OutputExportService,
    OutputSnapshotService,
)
from backend.app.risks import RiskRepository
from backend.app.schemas import RiskCreate, RiskEvidenceCreate, RiskTransition
from backend.tests.helpers import create_project
from backend.tests.test_risks import create_manual_risk, seed_evidence


def test_snapshot_requires_confirmed_risk_and_preserves_old_versions(tmp_path: Path) -> None:
    database = Database(tmp_path / "registry")
    project = create_project(database, tmp_path / "project")
    root = Path(project["storage_path"])
    document_id = seed_evidence(database, root)
    repository = RiskRepository(database)
    risk = create_manual_risk(database, project["id"], document_id)
    outputs = OutputSnapshotService(database)

    with pytest.raises(OutputError) as pending:
        outputs.create_risk_register(project["id"])
    assert pending.value.code == "NO_CONFIRMED_RISKS"

    verified = repository.transition(
        project["id"],
        risk["id"],
        RiskTransition(status="已核实", note="已逐页核对合成证据"),
    )
    first = outputs.create_risk_register(project["id"])
    first_risk = first["snapshot"]["risks"][0]

    assert first["risk_count"] == 1
    assert len(first["content_sha256"]) == 64
    assert first_risk["risk_version"] == verified["version"]
    assert first_risk["status"] == "已核实"
    assert [item["citation"] for item in first_risk["evidence"]] == [
        "R-0001-E01",
        "R-0001-E02",
    ]
    assert first_risk["evidence"][0]["source_reference"].endswith("第 1 页 · 块 1")

    repository.transition(
        project["id"],
        risk["id"],
        RiskTransition(status="已关闭", note="复核工作已完成"),
    )
    second = outputs.create_risk_register(project["id"])

    assert second["id"] != first["id"]
    assert second["snapshot"]["risks"][0]["status"] == "已关闭"
    assert outputs.get(project["id"], first["id"])["snapshot"]["risks"][0][
        "status"
    ] == "已核实"
    assert [item["id"] for item in outputs.list(project["id"])] == [
        second["id"],
        first["id"],
    ]

    with database.connect(root / "app.db") as db:
        events = db.execute(
            "SELECT details_json FROM audit_events WHERE event_type='output.snapshot_created'"
        ).fetchall()
    assert len(events) == 2
    assert json.loads(events[0]["details_json"])["risk_count"] == 1


def test_snapshot_rejects_confirmed_risk_without_support_evidence(tmp_path: Path) -> None:
    database = Database(tmp_path / "registry")
    project = create_project(database, tmp_path / "project")
    document_id = seed_evidence(database, Path(project["storage_path"]))
    repository = RiskRepository(database)
    risk = repository.create(
        project["id"],
        RiskCreate(
            summary="Counter evidence alone is insufficient for output",
            evidence=[
                RiskEvidenceCreate(
                    document_id=document_id,
                    page_number=2,
                    block_number=1,
                    quote="synthetic counter evidence text",
                    direction="counter",
                )
            ],
        ),
    )
    repository.transition(
        project["id"],
        risk["id"],
        RiskTransition(status="已核实", note="人工确认仍需补充支持证据"),
    )

    with pytest.raises(OutputError) as captured:
        OutputSnapshotService(database).create_risk_register(project["id"])

    assert captured.value.code == "RISK_SUPPORT_EVIDENCE_REQUIRED"
    assert OutputSnapshotService(database).list(project["id"]) == []


def test_output_draft_versions_finalize_and_export_without_overwrite(tmp_path: Path) -> None:
    database = Database(tmp_path / "registry")
    project = create_project(database, tmp_path / "project")
    root = Path(project["storage_path"])
    document_id = seed_evidence(database, root)
    repository = RiskRepository(database)
    risk = create_manual_risk(database, project["id"], document_id)
    repository.transition(
        project["id"],
        risk["id"],
        RiskTransition(status="已核实", note="已逐页核对合成证据"),
    )
    snapshots = OutputSnapshotService(database)
    snapshot = snapshots.create_risk_register(project["id"])
    drafts = OutputDraftService(database, snapshots)
    exports = OutputExportService(database, snapshots, drafts)

    draft = drafts.create(project["id"], snapshot["id"])
    assert draft["status"] == "editing"
    assert draft["version"] == 1
    assert draft["items"][0]["risk_number"] == "R-0001"
    assert draft["materials_title"] == "资料清单"
    assert draft["materials"] == [
        {
            "id": "M-001",
            "risk_id": risk["id"],
            "risk_number": "R-0001",
            "title": "R-0001 原始文件、审批记录及补充支持材料",
            "purpose": "用于复核“Synthetic evidence needs human review”的事实背景、期间归属和证据完整性。",
            "requested_scope": f"{project['entity_name']} · {project['year_start']}—{project['year_end']}",
            "priority": "待评估",
        }
    ]
    assert [item["id"] for item in draft["interviews"]] == ["Q-001", "Q-002"]

    with pytest.raises(OutputError) as pending:
        exports.create_excel(project["id"], draft["id"])
    assert pending.value.code == "OUTPUT_DRAFT_NOT_FINALIZED"
    with pytest.raises(OutputError) as pending_word:
        exports.create_word(project["id"], draft["id"])
    assert pending_word.value.code == "OUTPUT_DRAFT_NOT_FINALIZED"
    with pytest.raises(OutputError) as pending_pdf:
        exports.create_pdf(project["id"], draft["id"])
    assert pending_pdf.value.code == "OUTPUT_DRAFT_NOT_FINALIZED"

    updated = drafts.update(
        project["id"],
        draft["id"],
        title="合成审计风险清单（复核稿）",
        notes="仅用于离线验收。",
        items=[
            {
                "risk_id": draft["items"][0]["risk_id"],
                "heading": "银行存款余额异常需补充复核",
                "body": "已核对回函，并记录后续程序。",
            }
        ],
        materials_title="复核资料清单",
        materials=[
            {
                "id": draft["materials"][0]["id"],
                "risk_id": draft["materials"][0]["risk_id"],
                "title": "银行回函及期后流水",
                "purpose": draft["materials"][0]["purpose"],
                "requested_scope": draft["materials"][0]["requested_scope"],
                "priority": "高",
            }
        ],
        interview_title="复核访谈提纲",
        interviews=[
            {
                "id": item["id"],
                "risk_id": item["risk_id"],
                "audience": item["audience"],
                "question": item["question"],
                "objective": item["objective"],
            }
            for item in reversed(draft["interviews"])
        ],
    )
    assert updated["version"] == 2
    assert updated["materials"][0]["title"] == "银行回函及期后流水"
    assert [item["id"] for item in updated["interviews"]] == ["Q-002", "Q-001"]
    assert [item["version"] for item in updated["versions"]] == [2, 1]

    finalized = drafts.finalize(project["id"], draft["id"])
    assert finalized["status"] == "finalized"
    assert finalized["version"] == 3
    assert finalized["finalized_at"] is not None
    with pytest.raises(OutputError) as locked:
        drafts.update(
            project["id"],
            draft["id"],
            title="不应写入",
            notes="",
            items=updated["items"],
            materials_title=updated["materials_title"],
            materials=updated["materials"],
            interview_title=updated["interview_title"],
            interviews=updated["interviews"],
        )
    assert locked.value.code == "OUTPUT_DRAFT_FINALIZED"

    first = exports.create_excel(project["id"], draft["id"])
    second = exports.create_excel(project["id"], draft["id"])
    assert first["id"] != second["id"]
    assert first["filename"] != second["filename"]
    assert len(first["file_sha256"]) == 64
    assert first["draft_version"] == 3
    first_record, first_path = exports.get(project["id"], first["id"])
    second_record, second_path = exports.get(project["id"], second["id"])
    assert first_record["download_url"].endswith(f"/{first['id']}/file")
    assert first_path != second_path
    assert first_path.is_file() and second_path.is_file()
    assert len(exports.list(project["id"], draft["id"])) == 2

    workbook = load_workbook(first_path, read_only=True)
    risk_sheet, rule_sheet, evidence_sheet = workbook.worksheets
    assert risk_sheet["A1"].value == "合成审计风险清单（复核稿）"
    assert risk_sheet["B8"].value == "银行存款余额异常需补充复核"
    assert risk_sheet["G8"].value == "已核对回函，并记录后续程序。"
    assert rule_sheet["A8"].value == "R-0001"
    assert evidence_sheet["B8"].value == "R-0001-E01"
    workbook.close()

    first_word = exports.create_word(project["id"], draft["id"])
    second_word = exports.create_word(project["id"], draft["id"])
    assert first_word["id"] != second_word["id"]
    assert first_word["filename"] != second_word["filename"]
    assert first_word["export_format"] == "docx"
    assert first_word["template_version"] == "audit-work-products-word-v1"
    _, first_word_path = exports.get(project["id"], first_word["id"])
    _, second_word_path = exports.get(project["id"], second_word["id"])
    assert first_word_path != second_word_path
    document = WordDocument(first_word_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "合成审计风险清单（复核稿）" in table_text
    assert "银行存款余额异常需补充复核" in text
    assert "R-0001-E01" in table_text
    assert "复核资料清单" in text
    assert "银行回函及期后流水" in table_text
    assert "复核访谈提纲" in text
    assert updated["interviews"][0]["question"] in text
    assert len(exports.list(project["id"], draft["id"])) == 4

    first_pdf = exports.create_pdf(project["id"], draft["id"])
    second_pdf = exports.create_pdf(project["id"], draft["id"])
    assert first_pdf["id"] != second_pdf["id"]
    assert first_pdf["filename"] != second_pdf["filename"]
    assert first_pdf["export_format"] == "pdf"
    assert first_pdf["template_version"] == "audit-work-products-pdf-v1"
    _, first_pdf_path = exports.get(project["id"], first_pdf["id"])
    _, second_pdf_path = exports.get(project["id"], second_pdf["id"])
    assert first_pdf_path != second_pdf_path
    reader = PdfReader(first_pdf_path)
    assert len(reader.pages) >= 4
    assert reader.metadata.title == "合成审计风险清单（复核稿）"
    embedded_fonts = []
    for page in reader.pages:
        for font_reference in page["/Resources"]["/Font"].values():
            font = font_reference.get_object()
            descriptor_reference = font.get("/FontDescriptor")
            if descriptor_reference is None:
                continue
            descriptor = descriptor_reference.get_object()
            embedded_fonts.append(
                any(key in descriptor for key in ("/FontFile", "/FontFile2", "/FontFile3"))
            )
    assert any(embedded_fonts)
    pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "合成审计风险清单（复核稿）" in pdf_text
    assert "银行存款余额异常需补充复核" in pdf_text
    assert "R-0001-E01" in pdf_text
    assert "复核资料清单" in pdf_text
    assert "银行回函及期后流水" in pdf_text
    assert "复核访谈提纲" in pdf_text
    assert updated["interviews"][0]["question"] in pdf_text
    assert len(exports.list(project["id"], draft["id"])) == 6

    with database.connect(root / "app.db") as db:
        db.execute("UPDATE output_drafts SET materials_json='[]' WHERE id=?", (draft["id"],))
    with pytest.raises(OutputError) as legacy_word:
        exports.create_word(project["id"], draft["id"])
    assert legacy_word.value.code == "OUTPUT_DRAFT_SECTIONS_REQUIRED"
    with pytest.raises(OutputError) as legacy_pdf:
        exports.create_pdf(project["id"], draft["id"])
    assert legacy_pdf.value.code == "OUTPUT_DRAFT_SECTIONS_REQUIRED"

    with database.connect(root / "app.db") as db:
        events = {
            row["event_type"]
            for row in db.execute(
                "SELECT event_type FROM audit_events WHERE event_type LIKE 'output.%'"
            ).fetchall()
        }
    assert {
        "output.snapshot_created",
        "output.draft_created",
        "output.draft_updated",
        "output.draft_finalized",
        "output.export_created",
    } <= events


def test_output_draft_rejects_risk_set_changes(tmp_path: Path) -> None:
    database = Database(tmp_path / "registry")
    project = create_project(database, tmp_path / "project")
    root = Path(project["storage_path"])
    document_id = seed_evidence(database, root)
    repository = RiskRepository(database)
    risk = create_manual_risk(database, project["id"], document_id)
    repository.transition(
        project["id"], risk["id"], RiskTransition(status="已核实", note="确认合成风险")
    )
    snapshots = OutputSnapshotService(database)
    snapshot = snapshots.create_risk_register(project["id"])
    drafts = OutputDraftService(database, snapshots)
    draft = drafts.create(project["id"], snapshot["id"])

    with pytest.raises(OutputError) as captured:
        drafts.update(
            project["id"],
            draft["id"],
            title=draft["title"],
            notes="",
            items=[{"risk_id": "other", "heading": "错误风险", "body": ""}],
            materials_title=draft["materials_title"],
            materials=draft["materials"],
            interview_title=draft["interview_title"],
            interviews=draft["interviews"],
        )

    assert captured.value.code == "OUTPUT_DRAFT_RISK_SET_CHANGED"
    assert drafts.get(project["id"], draft["id"])["version"] == 1


def test_output_draft_rejects_material_set_and_interview_risk_link_changes(
    tmp_path: Path,
) -> None:
    database = Database(tmp_path / "registry")
    project = create_project(database, tmp_path / "project")
    document_id = seed_evidence(database, Path(project["storage_path"]))
    repository = RiskRepository(database)
    risk = create_manual_risk(database, project["id"], document_id)
    repository.transition(
        project["id"], risk["id"], RiskTransition(status="已核实", note="确认合成风险")
    )
    snapshots = OutputSnapshotService(database)
    draft = OutputDraftService(database, snapshots).create(
        project["id"], snapshots.create_risk_register(project["id"])["id"]
    )
    drafts = OutputDraftService(database, snapshots)
    common = {
        "title": draft["title"],
        "notes": draft["notes"],
        "items": draft["items"],
        "materials_title": draft["materials_title"],
        "interview_title": draft["interview_title"],
    }

    with pytest.raises(OutputError) as removed:
        drafts.update(
            project["id"],
            draft["id"],
            **common,
            materials=[],
            interviews=draft["interviews"],
        )
    assert removed.value.code == "OUTPUT_DRAFT_ITEM_SET_CHANGED"

    changed = [dict(item) for item in draft["interviews"]]
    changed[0]["risk_id"] = "other-risk"
    with pytest.raises(OutputError) as relinked:
        drafts.update(
            project["id"],
            draft["id"],
            **common,
            materials=draft["materials"],
            interviews=changed,
        )
    assert relinked.value.code == "OUTPUT_DRAFT_ITEM_LINK_CHANGED"


@pytest.mark.parametrize("value", ["=1+1", "+cmd", "-2+3", "@SUM(A1:A2)", "\t=1+1"])
def test_excel_export_writes_untrusted_formula_prefixes_as_text(
    tmp_path: Path, value: str
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A7"] = "字段"
    sheet["A8"] = "模板"

    OutputExportService._write_rows(sheet, [[value]], "SafeTextTable")
    target = tmp_path / "safe.xlsx"
    workbook.save(target)
    workbook.close()

    reopened = load_workbook(target, data_only=False)
    cell = reopened.active["A8"]
    assert cell.value == value
    assert cell.data_type == "s"
    reopened.close()
