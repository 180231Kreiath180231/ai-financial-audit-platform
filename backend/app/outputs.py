from __future__ import annotations

import hashlib
import json
import os
import uuid
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

from .db import Database, utc_now
from .risks import RiskRepository

OUTPUT_SCHEMA_VERSION = "output-snapshot-v1"
RISK_REGISTER_TEMPLATE_VERSION = "format-neutral-risk-register-v1"
CONFIRMED_RISK_STATUSES = ("已核实", "已关闭")
EXCEL_TEMPLATE_VERSION = "risk-register-excel-v1"
EXCEL_TEMPLATE_PATH = (
    Path(__file__).resolve().parents[2] / "templates" / "risk-register-excel-v1.xlsx"
)


class OutputError(RuntimeError):
    def __init__(self, code: str, message: str, action: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.action = action


class OutputSnapshotService:
    """Create immutable, format-neutral snapshots from confirmed server-side risks."""

    def __init__(self, database: Database) -> None:
        self.database = database
        self.risks = RiskRepository(database)

    def list(self, project_id: str) -> list[dict[str, Any]]:
        root = self.database.project_root(project_id)
        with self.database.connect(root / "app.db") as db:
            rows = db.execute(
                """SELECT id, output_kind, schema_version, template_version,
                risk_count, content_sha256, created_at
                FROM output_snapshots ORDER BY created_at DESC, id DESC"""
            ).fetchall()
        return [dict(row) for row in rows]

    def get(self, project_id: str, snapshot_id: str) -> dict[str, Any]:
        root = self.database.project_root(project_id)
        with self.database.connect(root / "app.db") as db:
            row = db.execute(
                "SELECT * FROM output_snapshots WHERE id=?",
                (snapshot_id,),
            ).fetchone()
        if row is None:
            raise KeyError(snapshot_id)
        result = dict(row)
        result["snapshot"] = json.loads(result.pop("snapshot_json"))
        return result

    def create_risk_register(self, project_id: str) -> dict[str, Any]:
        project = self.database.get_project(project_id)
        root = self.database.project_root(project_id)
        snapshot_id = str(uuid.uuid4())
        created_at = utc_now()

        with self.database.connect(root / "app.db") as db:
            db.execute("BEGIN IMMEDIATE")
            try:
                placeholders = ",".join("?" for _ in CONFIRMED_RISK_STATUSES)
                rows = db.execute(
                    f"""SELECT * FROM risk_items WHERE status IN ({placeholders})
                    ORDER BY risk_number""",
                    CONFIRMED_RISK_STATUSES,
                ).fetchall()
                if not rows:
                    raise OutputError(
                        "NO_CONFIRMED_RISKS",
                        "当前项目没有可固化的已核实风险",
                        "先在风险台账完成证据复核并将风险转为“已核实”",
                    )

                risks = [
                    self.risks._hydrate(db, row, include_versions=False)  # noqa: SLF001
                    for row in rows
                ]
                snapshot_risks = [self._snapshot_risk(risk) for risk in risks]
                payload = {
                    "schema_version": OUTPUT_SCHEMA_VERSION,
                    "output_kind": "risk_register",
                    "project": {
                        "id": project["id"],
                        "name": project["name"],
                        "entity_name": project["entity_name"],
                        "year_start": project["year_start"],
                        "year_end": project["year_end"],
                    },
                    "risks": snapshot_risks,
                }
                serialized = json.dumps(
                    payload,
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
                content_sha256 = hashlib.sha256(serialized.encode("utf-8")).hexdigest()
                db.execute(
                    """INSERT INTO output_snapshots
                    (id, output_kind, schema_version, template_version, risk_count,
                     content_sha256, snapshot_json, created_at)
                    VALUES (?, 'risk_register', ?, ?, ?, ?, ?, ?)""",
                    (
                        snapshot_id,
                        OUTPUT_SCHEMA_VERSION,
                        RISK_REGISTER_TEMPLATE_VERSION,
                        len(snapshot_risks),
                        content_sha256,
                        serialized,
                        created_at,
                    ),
                )
                db.executemany(
                    """INSERT INTO output_snapshot_risks
                    (snapshot_id, risk_id, risk_number, risk_version, risk_status)
                    VALUES (?, ?, ?, ?, ?)""",
                    [
                        (
                            snapshot_id,
                            risk["risk_id"],
                            risk["risk_number"],
                            risk["risk_version"],
                            risk["status"],
                        )
                        for risk in snapshot_risks
                    ],
                )
                db.execute(
                    "INSERT INTO audit_events VALUES (?, NULL, ?, ?, ?)",
                    (
                        str(uuid.uuid4()),
                        "output.snapshot_created",
                        created_at,
                        json.dumps(
                            {
                                "snapshot_id": snapshot_id,
                                "output_kind": "risk_register",
                                "risk_count": len(snapshot_risks),
                                "content_sha256": content_sha256,
                            },
                            ensure_ascii=False,
                        ),
                    ),
                )
                db.execute("COMMIT")
            except Exception:
                db.execute("ROLLBACK")
                raise
        return self.get(project_id, snapshot_id)

    @staticmethod
    def _snapshot_risk(risk: dict[str, Any]) -> dict[str, Any]:
        evidence = []
        for index, item in enumerate(risk["evidence"], start=1):
            if item["kind"] == "document":
                source_reference = (
                    f"{item['document_name']} · 第 {item['page_number']} 页"
                    f" · 块 {item['block_number']}"
                )
            else:
                source_reference = (
                    f"{item['document_name']} · CSV 行 {item['line_start']}–{item['line_end']}"
                )
            evidence.append(
                {
                    "citation": f"{risk['risk_number']}-E{index:02d}",
                    "source_evidence_id": item["id"],
                    "kind": item["kind"],
                    "direction": item["direction"],
                    "source_reference": source_reference,
                    "quote": item["quote"],
                    "document_id": item.get("document_id"),
                    "page_number": item.get("page_number"),
                    "block_number": item.get("block_number"),
                    "dataset_id": item.get("dataset_id"),
                    "line_start": item.get("line_start"),
                    "line_end": item.get("line_end"),
                    "period_key": item.get("period_key"),
                    "account_code": item.get("account_code"),
                    "parse_method": item.get("parse_method"),
                    "parse_version": item.get("parse_version"),
                }
            )
        if not any(item["direction"] == "support" for item in evidence):
            raise OutputError(
                "RISK_SUPPORT_EVIDENCE_REQUIRED",
                f"{risk['risk_number']} 缺少支持证据，不能生成快照",
                "回到风险台账补充并核对支持证据",
            )
        return {
            "risk_id": risk["id"],
            "risk_number": risk["risk_number"],
            "risk_version": risk["version"],
            "risk_type": risk["risk_type"],
            "risk_level": risk["risk_level"],
            "status": risk["status"],
            "summary": risk["summary"],
            "trigger_rule_id": risk["trigger_rule_id"],
            "trigger_rule_version": risk["trigger_rule_version"],
            "input_values": risk["input_values"],
            "baseline_values": risk["baseline_values"],
            "calculation_result": risk["calculation_result"],
            "model_explanation": risk["model_explanation"],
            "uncertainty": risk["uncertainty"],
            "human_opinion": risk["human_opinion"],
            "model_source": {
                "provider": risk["model_provider"],
                "actual_model": risk["actual_model"],
                "model_call_id": risk["model_call_id"],
            },
            "evidence": evidence,
        }


class OutputDraftService:
    """Version editable output copy without mutating its source snapshot."""

    def __init__(self, database: Database, snapshots: OutputSnapshotService) -> None:
        self.database = database
        self.snapshots = snapshots

    def list(self, project_id: str, snapshot_id: str | None = None) -> list[dict[str, Any]]:
        root = self.database.project_root(project_id)
        query = "SELECT * FROM output_drafts"
        parameters: tuple[str, ...] = ()
        if snapshot_id:
            query += " WHERE snapshot_id=?"
            parameters = (snapshot_id,)
        query += " ORDER BY created_at DESC, id DESC"
        with self.database.connect(root / "app.db") as db:
            rows = db.execute(query, parameters).fetchall()
            return [self._hydrate(db, row) for row in rows]

    def get(self, project_id: str, draft_id: str) -> dict[str, Any]:
        root = self.database.project_root(project_id)
        with self.database.connect(root / "app.db") as db:
            row = db.execute("SELECT * FROM output_drafts WHERE id=?", (draft_id,)).fetchone()
            if row is None:
                raise KeyError(draft_id)
            return self._hydrate(db, row)

    def create(self, project_id: str, snapshot_id: str) -> dict[str, Any]:
        snapshot = self.snapshots.get(project_id, snapshot_id)
        root = self.database.project_root(project_id)
        now = utc_now()
        draft_id = str(uuid.uuid4())
        items = [
            {
                "risk_id": risk["risk_id"],
                "risk_number": risk["risk_number"],
                "heading": risk["summary"],
                "body": risk["human_opinion"] or risk["model_explanation"] or "",
            }
            for risk in snapshot["snapshot"]["risks"]
        ]
        title = f"{snapshot['snapshot']['project']['name']}风险清单"
        materials = self._build_materials(snapshot["snapshot"])
        interviews = self._build_interviews(snapshot["snapshot"])
        materials_title = "资料清单"
        interview_title = "访谈提纲"
        serialized_items = self._serialize(items)
        draft_state = self._version_state(
            status="editing",
            title=title,
            notes="",
            items=items,
            materials_title=materials_title,
            materials=materials,
            interview_title=interview_title,
            interviews=interviews,
        )
        with self.database.connect(root / "app.db") as db:
            db.execute("BEGIN IMMEDIATE")
            try:
                db.execute(
                    """INSERT INTO output_drafts
                    (id, snapshot_id, output_kind, status, version, title, notes, items_json,
                     created_at, updated_at, finalized_at, materials_title, materials_json,
                     interview_title, interview_json)
                    VALUES (?, ?, 'risk_register', 'editing', 1, ?, '', ?, ?, ?, NULL,
                            ?, ?, ?, ?)""",
                    (
                        draft_id,
                        snapshot_id,
                        title,
                        serialized_items,
                        now,
                        now,
                        materials_title,
                        self._serialize(materials),
                        interview_title,
                        self._serialize(interviews),
                    ),
                )
                self._insert_version(db, draft_id, 1, "从不可变快照创建草稿", draft_state, now)
                self._audit(
                    db,
                    "output.draft_created",
                    now,
                    {"draft_id": draft_id, "snapshot_id": snapshot_id, "version": 1},
                )
                db.execute("COMMIT")
            except Exception:
                db.execute("ROLLBACK")
                raise
        return self.get(project_id, draft_id)

    def update(
        self,
        project_id: str,
        draft_id: str,
        *,
        title: str,
        notes: str,
        items: list[dict[str, str]],
        materials_title: str,
        materials: list[dict[str, str]],
        interview_title: str,
        interviews: list[dict[str, str]],
    ) -> dict[str, Any]:
        root = self.database.project_root(project_id)
        now = utc_now()
        with self.database.connect(root / "app.db") as db:
            db.execute("BEGIN IMMEDIATE")
            try:
                row = db.execute("SELECT * FROM output_drafts WHERE id=?", (draft_id,)).fetchone()
                if row is None:
                    raise KeyError(draft_id)
                if row["status"] != "editing":
                    raise OutputError(
                        "OUTPUT_DRAFT_FINALIZED",
                        "已最终固化的草稿不能继续修改",
                        "从原始快照创建一份新草稿后再编辑",
                    )
                current_items = json.loads(row["items_json"])
                expected = {item["risk_id"] for item in current_items}
                received = [item["risk_id"] for item in items]
                if len(received) != len(set(received)) or set(received) != expected:
                    raise OutputError(
                        "OUTPUT_DRAFT_RISK_SET_CHANGED",
                        "草稿必须保留快照中的全部风险，且每项只能出现一次",
                        "仅调整顺序或编辑标题和正文，不要新增、删除或重复风险",
                    )
                risk_numbers = {item["risk_id"]: item["risk_number"] for item in current_items}
                normalized = [
                    {
                        "risk_id": item["risk_id"],
                        "risk_number": risk_numbers[item["risk_id"]],
                        "heading": item["heading"],
                        "body": item["body"],
                    }
                    for item in items
                ]
                current_materials = json.loads(row["materials_json"])
                normalized_materials = self._normalize_linked_items(
                    current_materials,
                    materials,
                    kind="资料清单",
                    editable_fields=("title", "purpose", "requested_scope", "priority"),
                )
                current_interviews = json.loads(row["interview_json"])
                normalized_interviews = self._normalize_linked_items(
                    current_interviews,
                    interviews,
                    kind="访谈提纲",
                    editable_fields=("audience", "question", "objective"),
                )
                if (
                    title == row["title"]
                    and notes == row["notes"]
                    and normalized == current_items
                    and materials_title == row["materials_title"]
                    and normalized_materials == current_materials
                    and interview_title == row["interview_title"]
                    and normalized_interviews == current_interviews
                ):
                    db.execute("ROLLBACK")
                    return self.get(project_id, draft_id)
                version = row["version"] + 1
                db.execute(
                    """UPDATE output_drafts SET version=?, title=?, notes=?, items_json=?,
                    materials_title=?, materials_json=?, interview_title=?, interview_json=?,
                    updated_at=? WHERE id=?""",
                    (
                        version,
                        title,
                        notes,
                        self._serialize(normalized),
                        materials_title,
                        self._serialize(normalized_materials),
                        interview_title,
                        self._serialize(normalized_interviews),
                        now,
                        draft_id,
                    ),
                )
                state = self._version_state(
                    status="editing",
                    title=title,
                    notes=notes,
                    items=normalized,
                    materials_title=materials_title,
                    materials=normalized_materials,
                    interview_title=interview_title,
                    interviews=normalized_interviews,
                )
                self._insert_version(db, draft_id, version, "人工编辑输出草稿", state, now)
                self._audit(
                    db,
                    "output.draft_updated",
                    now,
                    {"draft_id": draft_id, "version": version},
                )
                db.execute("COMMIT")
            except Exception:
                if db.in_transaction:
                    db.execute("ROLLBACK")
                raise
        return self.get(project_id, draft_id)

    def finalize(self, project_id: str, draft_id: str) -> dict[str, Any]:
        root = self.database.project_root(project_id)
        now = utc_now()
        with self.database.connect(root / "app.db") as db:
            db.execute("BEGIN IMMEDIATE")
            try:
                row = db.execute("SELECT * FROM output_drafts WHERE id=?", (draft_id,)).fetchone()
                if row is None:
                    raise KeyError(draft_id)
                if row["status"] == "finalized":
                    db.execute("ROLLBACK")
                    return self.get(project_id, draft_id)
                version = row["version"] + 1
                db.execute(
                    """UPDATE output_drafts SET status='finalized', version=?, updated_at=?,
                    finalized_at=? WHERE id=?""",
                    (version, now, now, draft_id),
                )
                state = self._version_state(
                    status="finalized",
                    title=row["title"],
                    notes=row["notes"],
                    items=json.loads(row["items_json"]),
                    materials_title=row["materials_title"],
                    materials=json.loads(row["materials_json"]),
                    interview_title=row["interview_title"],
                    interviews=json.loads(row["interview_json"]),
                )
                self._insert_version(db, draft_id, version, "最终固化输出草稿", state, now)
                self._audit(
                    db,
                    "output.draft_finalized",
                    now,
                    {"draft_id": draft_id, "version": version},
                )
                db.execute("COMMIT")
            except Exception:
                if db.in_transaction:
                    db.execute("ROLLBACK")
                raise
        return self.get(project_id, draft_id)

    @staticmethod
    def _hydrate(db, row) -> dict[str, Any]:
        result = dict(row)
        result["items"] = json.loads(result.pop("items_json"))
        result["materials"] = json.loads(result.pop("materials_json"))
        result["interviews"] = json.loads(result.pop("interview_json"))
        versions = db.execute(
            """SELECT version, change_reason, created_at FROM output_draft_versions
            WHERE draft_id=? ORDER BY version DESC""",
            (result["id"],),
        ).fetchall()
        result["versions"] = [dict(version) for version in versions]
        return result

    @staticmethod
    def _serialize(value: Any) -> str:
        return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))

    @staticmethod
    def _version_state(
        *,
        status: str,
        title: str,
        notes: str,
        items: list[dict[str, Any]],
        materials_title: str,
        materials: list[dict[str, Any]],
        interview_title: str,
        interviews: list[dict[str, Any]],
    ) -> dict[str, Any]:
        return {
            "status": status,
            "title": title,
            "notes": notes,
            "items": items,
            "materials_title": materials_title,
            "materials": materials,
            "interview_title": interview_title,
            "interviews": interviews,
        }

    @staticmethod
    def _build_materials(snapshot: dict[str, Any]) -> list[dict[str, str]]:
        project = snapshot["project"]
        requested_scope = (
            f"{project['entity_name']} · {project['year_start']}—{project['year_end']}"
        )
        priorities = {"高": "高", "中": "中", "低": "低"}
        result: list[dict[str, str]] = []
        for index, risk in enumerate(snapshot["risks"], start=1):
            has_financial_source = any(
                evidence["kind"] == "financial" for evidence in risk["evidence"]
            )
            title = (
                f"{risk['risk_number']} 科目明细、余额形成依据及支持性资料"
                if has_financial_source
                else f"{risk['risk_number']} 原始文件、审批记录及补充支持材料"
            )
            result.append(
                {
                    "id": f"M-{index:03d}",
                    "risk_id": risk["risk_id"],
                    "risk_number": risk["risk_number"],
                    "title": title,
                    "purpose": (
                        f"用于复核“{risk['summary']}”的事实背景、期间归属和证据完整性。"
                    ),
                    "requested_scope": requested_scope,
                    "priority": priorities.get(risk["risk_level"], "待评估"),
                }
            )
        return result

    @staticmethod
    def _build_interviews(snapshot: dict[str, Any]) -> list[dict[str, str]]:
        result: list[dict[str, str]] = []
        for risk in snapshot["risks"]:
            number = risk["risk_number"]
            summary = risk["summary"]
            sequence = len(result) + 1
            result.extend(
                [
                    {
                        "id": f"Q-{sequence:03d}",
                        "risk_id": risk["risk_id"],
                        "risk_number": number,
                        "audience": "业务负责人 / 财务负责人",
                        "question": f"请说明“{summary}”的业务背景、发生原因和涉及期间。",
                        "objective": f"核实 {number} 的事实背景和时间范围。",
                    },
                    {
                        "id": f"Q-{sequence + 1:03d}",
                        "risk_id": risk["risk_id"],
                        "risk_number": number,
                        "audience": "业务负责人 / 财务负责人",
                        "question": (
                            f"围绕 {number}，相关支持材料如何形成、由谁复核，"
                            "是否存在未记录的例外或反证？"
                        ),
                        "objective": f"了解 {number} 的证据形成、复核责任和潜在反证。",
                    },
                ]
            )
        return result

    @staticmethod
    def _normalize_linked_items(
        current: list[dict[str, Any]],
        received: list[dict[str, Any]],
        *,
        kind: str,
        editable_fields: tuple[str, ...],
    ) -> list[dict[str, Any]]:
        current_by_id = {item["id"]: item for item in current}
        received_ids = [item["id"] for item in received]
        if len(received_ids) != len(set(received_ids)) or set(received_ids) != set(
            current_by_id
        ):
            raise OutputError(
                "OUTPUT_DRAFT_ITEM_SET_CHANGED",
                f"{kind}必须保留系统生成的全部项目，且每项只能出现一次",
                "仅调整顺序或编辑文案，不要新增、删除或重复项目",
            )
        normalized = []
        for item in received:
            source = current_by_id[item["id"]]
            if item["risk_id"] != source["risk_id"]:
                raise OutputError(
                    "OUTPUT_DRAFT_ITEM_LINK_CHANGED",
                    f"{kind}项目不能改绑到其他风险",
                    "保留项目与风险的原始关联，只编辑允许修改的内容",
                )
            normalized.append(
                {
                    "id": source["id"],
                    "risk_id": source["risk_id"],
                    "risk_number": source["risk_number"],
                    **{field: item[field] for field in editable_fields},
                }
            )
        return normalized

    def _insert_version(
        self,
        db,
        draft_id: str,
        version: int,
        reason: str,
        state: dict[str, Any],
        created_at: str,
    ) -> None:
        db.execute(
            "INSERT INTO output_draft_versions VALUES (?, ?, ?, ?, ?)",
            (draft_id, version, reason, self._serialize(state), created_at),
        )

    @staticmethod
    def _audit(db, event_type: str, created_at: str, details: dict[str, Any]) -> None:
        db.execute(
            "INSERT INTO audit_events VALUES (?, NULL, ?, ?, ?)",
            (str(uuid.uuid4()), event_type, created_at, json.dumps(details, ensure_ascii=False)),
        )


class OutputExportService:
    """Render immutable Excel exports from finalized draft versions."""

    def __init__(
        self,
        database: Database,
        snapshots: OutputSnapshotService,
        drafts: OutputDraftService,
    ) -> None:
        self.database = database
        self.snapshots = snapshots
        self.drafts = drafts

    def list(self, project_id: str, draft_id: str | None = None) -> list[dict[str, Any]]:
        root = self.database.project_root(project_id)
        query = "SELECT * FROM output_exports"
        parameters: tuple[str, ...] = ()
        if draft_id:
            query += " WHERE draft_id=?"
            parameters = (draft_id,)
        query += " ORDER BY created_at DESC, id DESC"
        with self.database.connect(root / "app.db") as db:
            rows = db.execute(query, parameters).fetchall()
        return [self._public_record(project_id, dict(row)) for row in rows]

    def create_excel(self, project_id: str, draft_id: str) -> dict[str, Any]:
        draft = self.drafts.get(project_id, draft_id)
        if draft["status"] != "finalized":
            raise OutputError(
                "OUTPUT_DRAFT_NOT_FINALIZED",
                "草稿尚未最终固化，不能导出",
                "确认标题、正文、顺序和备注后先执行最终固化",
            )
        snapshot = self.snapshots.get(project_id, draft["snapshot_id"])
        if not EXCEL_TEMPLATE_PATH.is_file():
            raise OutputError(
                "OUTPUT_TEMPLATE_MISSING",
                "Excel 模板文件不存在",
                "恢复仓库 templates/risk-register-excel-v1.xlsx 后重试",
            )
        root = self.database.project_root(project_id)
        exports_dir = (root / "exports").resolve()
        exports_dir.mkdir(parents=True, exist_ok=True)
        export_id = str(uuid.uuid4())
        created_at = utc_now()
        date_token = created_at[:10].replace("-", "")
        filename = f"风险清单-{date_token}-v{draft['version']}-{export_id[:8]}.xlsx"
        final_path = (exports_dir / f"{export_id}.xlsx").resolve()
        temp_path = (root / "temp" / f"{export_id}.xlsx.tmp").resolve()
        temp_path.parent.mkdir(parents=True, exist_ok=True)
        if exports_dir not in final_path.parents:
            raise OutputError("OUTPUT_PATH_INVALID", "导出路径无效", "检查项目存储目录后重试")
        try:
            self._render_excel(snapshot, draft, temp_path, created_at)
            file_sha256 = self._sha256(temp_path)
            size_bytes = temp_path.stat().st_size
            os.replace(temp_path, final_path)
            relative_path = final_path.relative_to(root.resolve()).as_posix()
            with self.database.connect(root / "app.db") as db:
                db.execute("BEGIN IMMEDIATE")
                try:
                    db.execute(
                        """INSERT INTO output_exports
                        (id, draft_id, draft_version, snapshot_id, export_format,
                         template_version, filename, stored_path, file_sha256,
                         size_bytes, created_at)
                        VALUES (?, ?, ?, ?, 'xlsx', ?, ?, ?, ?, ?, ?)""",
                        (
                            export_id,
                            draft_id,
                            draft["version"],
                            draft["snapshot_id"],
                            EXCEL_TEMPLATE_VERSION,
                            filename,
                            relative_path,
                            file_sha256,
                            size_bytes,
                            created_at,
                        ),
                    )
                    OutputDraftService._audit(
                        db,
                        "output.export_created",
                        created_at,
                        {
                            "export_id": export_id,
                            "draft_id": draft_id,
                            "draft_version": draft["version"],
                            "snapshot_id": draft["snapshot_id"],
                            "format": "xlsx",
                            "template_version": EXCEL_TEMPLATE_VERSION,
                            "file_sha256": file_sha256,
                        },
                    )
                    db.execute("COMMIT")
                except Exception:
                    db.execute("ROLLBACK")
                    final_path.unlink(missing_ok=True)
                    raise
        finally:
            temp_path.unlink(missing_ok=True)
        return self.get(project_id, export_id)[0]

    def get(self, project_id: str, export_id: str) -> tuple[dict[str, Any], Path]:
        root = self.database.project_root(project_id).resolve()
        with self.database.connect(root / "app.db") as db:
            row = db.execute("SELECT * FROM output_exports WHERE id=?", (export_id,)).fetchone()
        if row is None:
            raise KeyError(export_id)
        record = dict(row)
        path = (root / record["stored_path"]).resolve()
        exports_dir = (root / "exports").resolve()
        if exports_dir not in path.parents or not path.is_file():
            raise OutputError(
                "OUTPUT_FILE_UNAVAILABLE",
                "导出文件不存在或路径无效",
                "重新从已固化草稿生成一份新导出",
            )
        return self._public_record(project_id, record), path

    @staticmethod
    def _public_record(project_id: str, record: dict[str, Any]) -> dict[str, Any]:
        record.pop("stored_path", None)
        record["download_url"] = (
            f"/api/v1/projects/{project_id}/outputs/exports/{record['id']}/file"
        )
        return record

    @staticmethod
    def _sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    def _render_excel(
        self,
        snapshot: dict[str, Any],
        draft: dict[str, Any],
        target: Path,
        created_at: str,
    ) -> None:
        workbook = load_workbook(EXCEL_TEMPLATE_PATH)
        risk_sheet, rule_sheet, evidence_sheet = workbook.worksheets
        project = snapshot["snapshot"]["project"]
        snapshot_risks = {
            risk["risk_id"]: risk for risk in snapshot["snapshot"]["risks"]
        }
        ordered = [(item, snapshot_risks[item["risk_id"]]) for item in draft["items"]]
        evidence_count = sum(len(risk["evidence"]) for _, risk in ordered)
        for sheet, count in (
            (risk_sheet, len(ordered)),
            (rule_sheet, len(ordered)),
            (evidence_sheet, evidence_count),
        ):
            sheet["B2"] = project["name"]
            sheet["E2"] = EXCEL_TEMPLATE_VERSION
            sheet["B3"] = project["entity_name"]
            sheet["E3"] = f"{project['year_start']}—{project['year_end']}"
            sheet["B4"] = snapshot["id"]
            sheet["E4"] = created_at
            sheet["B5"] = snapshot["content_sha256"]
            sheet["E5"] = str(count)
            sheet["E5"].alignment = Alignment(
                horizontal="left", vertical="center", wrap_text=True
            )
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.print_title_rows = "1:7"
        risk_sheet["A1"] = draft["title"]
        workbook.properties.title = draft["title"]
        workbook.properties.subject = (
            f"不可变快照 {snapshot['id']} / 最终草稿 {draft['id']} v{draft['version']}"
        )
        workbook.properties.keywords = snapshot["content_sha256"]
        risk_rows = []
        rule_rows = []
        evidence_rows = []
        for item, risk in ordered:
            supports = [e["citation"] for e in risk["evidence"] if e["direction"] == "support"]
            counters = [e["citation"] for e in risk["evidence"] if e["direction"] == "counter"]
            source = risk["model_source"]
            model_source = "未使用模型"
            if source["provider"] or source["actual_model"] or source["model_call_id"]:
                model_source = " / ".join(
                    value
                    for value in (
                        source["provider"],
                        source["actual_model"],
                        source["model_call_id"],
                    )
                    if value
                )
            risk_rows.append(
                [
                    risk["risk_number"], item["heading"], risk["risk_type"],
                    risk["risk_level"], risk["status"],
                    f"{risk['trigger_rule_id']} / {risk['trigger_rule_version']}",
                    item["body"], risk["uncertainty"], risk["model_explanation"] or "",
                    model_source, "、".join(supports), "、".join(counters),
                    risk["risk_version"],
                ]
            )
            rule_rows.append(
                [
                    risk["risk_number"],
                    f"{risk['trigger_rule_id']} / {risk['trigger_rule_version']}",
                    self._json_cell(risk["input_values"]),
                    self._json_cell(risk["baseline_values"]),
                    self._json_cell(risk["calculation_result"]),
                    risk["risk_version"],
                ]
            )
            for evidence in risk["evidence"]:
                evidence_rows.append(
                    [
                        risk["risk_number"], evidence["citation"],
                        "支持" if evidence["direction"] == "support" else "反证",
                        evidence["kind"], evidence["source_reference"], evidence["quote"],
                        evidence["parse_method"] or "", evidence["parse_version"] or "",
                        evidence["source_evidence_id"],
                    ]
                )
        self._write_rows(risk_sheet, risk_rows, "RiskRegisterTable")
        self._write_rows(rule_sheet, rule_rows, "RuleDetailsTable")
        self._write_rows(evidence_sheet, evidence_rows, "EvidenceIndexTable")
        workbook.save(target)

    @staticmethod
    def _json_cell(value: dict[str, Any]) -> str:
        return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(", ", ": "))

    @staticmethod
    def _write_rows(sheet, rows: list[list[Any]], table_name: str) -> None:
        template_row = 8
        for row_index, values in enumerate(rows, start=template_row):
            for column_index, value in enumerate(values, start=1):
                source = sheet.cell(template_row, column_index)
                cell = sheet.cell(row_index, column_index, value=value)
                if row_index != template_row:
                    cell._style = copy(source._style)  # noqa: SLF001
                    cell.font = copy(source.font)
                    cell.fill = copy(source.fill)
                    cell.border = copy(source.border)
                    cell.alignment = copy(source.alignment)
                    cell.number_format = source.number_format
                    cell.protection = copy(source.protection)
            sheet.row_dimensions[row_index].height = 42
        end_row = template_row + max(len(rows), 1) - 1
        end_column = sheet.max_column
        sheet.auto_filter.ref = f"A7:{sheet.cell(7, end_column).coordinate}"
        table = Table(displayName=table_name, ref=f"A7:{sheet.cell(end_row, end_column).coordinate}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
