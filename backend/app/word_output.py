from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

WORD_TEMPLATE_VERSION = "audit-work-products-word-v2"
WORD_TEMPLATE_PATH = (
    Path(__file__).resolve().parents[2] / "templates" / "audit-work-products-word-v2.docx"
)

_NAVY = "17365D"
_PALE_BLUE = "EAF1F8"
_PALE_GRAY = "F5F7F9"
_BORDER = "D9D9D9"
_MUTED = RGBColor(78, 91, 108)


def render_word_export(
    snapshot: dict[str, Any],
    draft: dict[str, Any],
    target: Path,
    created_at: str,
) -> None:
    document = Document(WORD_TEMPLATE_PATH)
    _remove_template_marker(document)
    payload = snapshot["snapshot"]
    project = payload["project"]
    risks = {risk["risk_id"]: risk for risk in payload["risks"]}

    document.core_properties.title = f"{project['name']} 审计工作成果"
    document.core_properties.subject = (
        f"不可变快照 {snapshot['id']} / 最终草稿 {draft['id']} v{draft['version']}"
    )
    document.core_properties.author = "衡鉴审计工作台"
    document.core_properties.keywords = snapshot["content_sha256"]

    document.add_paragraph(f"{project['name']} 审计工作成果", style="Title")
    subtitle = document.add_paragraph(
        "管理层沟通材料 风险清单 资料清单 访谈提纲", style="Subtitle"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    introduction = document.add_paragraph(
        "本文件汇集已最终固化的管理层沟通材料、风险清单、资料索取项目和访谈问题。"
        "风险事实、规则计算和证据引用来自不可变快照；管理层沟通内容用于确认事实和后续安排。"
    )
    introduction.paragraph_format.space_after = Pt(12)

    metadata = [
        ("成果包标题", draft["title"]),
        ("审计主体", project["entity_name"]),
        ("审计期间", f"{project['year_start']} 至 {project['year_end']}"),
        ("源快照", snapshot["id"]),
        ("最终草稿", f"v{draft['version']}"),
        ("生成时间", created_at),
    ]
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(1.25)
    table.columns[1].width = Inches(5.55)
    _set_cell(table.cell(0, 0), "字段", bold=True, fill=_NAVY, white=True, center=True)
    _set_cell(table.cell(0, 1), "内容", bold=True, fill=_NAVY, white=True, center=True)
    _repeat_header(table.rows[0])
    for label, value in metadata:
        cells = table.add_row().cells
        _set_cell(cells[0], label, bold=True, fill=_PALE_BLUE, center=False)
        _set_cell(cells[1], str(value), fill="FFFFFF", center=False)
    _format_table(table, font_size=10)

    counts = document.add_paragraph()
    counts.paragraph_format.space_before = Pt(10)
    counts.add_run("成果范围  ").bold = True
    counts.add_run(
        f"管理层事项 {len(draft['management'])} 项  ·  风险 {len(draft['items'])} 项  ·  "
        f"资料 {len(draft['materials'])} 项  ·  访谈问题 {len(draft['interviews'])} 项"
    )
    if draft["notes"]:
        notes = document.add_paragraph()
        notes.add_run("编制备注  ").bold = True
        notes.add_run(draft["notes"])

    document.add_page_break()
    document.add_heading(draft["management_title"], level=1)
    document.add_paragraph(
        "以下内容用于管理层沟通准备。风险状态、风险等级和证据编号保持源快照记录；"
        "管理层意见和后续安排应由相关人员确认后另行留痕。"
    )
    for item in draft["management"]:
        risk = risks[item["risk_id"]]
        document.add_heading(f"{risk['risk_number']} {item['heading']}", level=2)
        management_table = document.add_table(rows=1, cols=2)
        management_table.alignment = WD_TABLE_ALIGNMENT.LEFT
        management_table.autofit = False
        management_table.columns[0].width = Inches(1.25)
        management_table.columns[1].width = Inches(5.55)
        _set_cell(
            management_table.cell(0, 0),
            "字段",
            bold=True,
            fill=_NAVY,
            white=True,
            center=True,
        )
        _set_cell(
            management_table.cell(0, 1),
            "内容",
            bold=True,
            fill=_NAVY,
            white=True,
            center=True,
        )
        _repeat_header(management_table.rows[0])
        citations = "、".join(evidence["citation"] for evidence in risk["evidence"])
        for label, value in (
            ("风险等级", risk["risk_level"]),
            ("风险状态", risk["status"]),
            ("事项摘要", item["summary"]),
            ("需管理层回复", item["response_request"]),
            ("证据编号", citations),
        ):
            cells = management_table.add_row().cells
            _set_cell(cells[0], label, bold=True, fill=_PALE_GRAY)
            _set_cell(cells[1], str(value or ""))
        _format_table(management_table, font_size=9.5)

    document.add_page_break()
    document.add_heading("风险清单", level=1)
    document.add_paragraph(
        "以下内容按最终草稿顺序列示。风险编号、状态、规则、版本和证据引用保持源快照记录。"
    )
    for item in draft["items"]:
        risk = risks[item["risk_id"]]
        document.add_heading(f"{risk['risk_number']} {item['heading']}", level=2)
        if item["body"]:
            document.add_paragraph(item["body"])
        details = [
            ("风险类型", risk["risk_type"]),
            ("风险等级", risk["risk_level"]),
            ("状态", risk["status"]),
            ("触发规则", f"{risk['trigger_rule_id']} / {risk['trigger_rule_version']}"),
            ("风险版本", f"v{risk['risk_version']}"),
            ("不确定性", risk["uncertainty"]),
        ]
        detail_table = document.add_table(rows=1, cols=2)
        detail_table.alignment = WD_TABLE_ALIGNMENT.LEFT
        detail_table.autofit = False
        detail_table.columns[0].width = Inches(1.15)
        detail_table.columns[1].width = Inches(5.65)
        _set_cell(detail_table.cell(0, 0), "字段", bold=True, fill=_NAVY, white=True, center=True)
        _set_cell(detail_table.cell(0, 1), "内容", bold=True, fill=_NAVY, white=True, center=True)
        _repeat_header(detail_table.rows[0])
        for label, value in details:
            cells = detail_table.add_row().cells
            _set_cell(cells[0], label, bold=True, fill=_PALE_GRAY)
            _set_cell(cells[1], str(value or ""))
        _format_table(detail_table, font_size=9.5)
        document.add_heading("证据索引", level=3)
        evidence_table = document.add_table(rows=1, cols=4)
        evidence_table.alignment = WD_TABLE_ALIGNMENT.LEFT
        evidence_table.autofit = False
        for index, width in enumerate((0.8, 0.65, 2.0, 3.35)):
            evidence_table.columns[index].width = Inches(width)
        for index, label in enumerate(("编号", "方向", "来源定位", "原文引用")):
            _set_cell(evidence_table.cell(0, index), label, bold=True, fill=_NAVY, white=True, center=True)
        _repeat_header(evidence_table.rows[0])
        for row_index, evidence in enumerate(risk["evidence"], start=1):
            cells = evidence_table.add_row().cells
            values = (
                evidence["citation"],
                "支持" if evidence["direction"] == "support" else "反证",
                evidence["source_reference"],
                evidence["quote"],
            )
            for index, value in enumerate(values):
                _set_cell(
                    cells[index],
                    str(value),
                    fill=_PALE_BLUE if row_index % 2 == 0 else "FFFFFF",
                    center=index < 2,
                )
        _format_table(evidence_table, font_size=9)

    document.add_page_break()
    document.add_heading(draft["materials_title"], level=1)
    document.add_paragraph(
        "资料项目按最终草稿顺序列示。项目编号和关联风险保持锁定，具体提供方式与时间由审计人员另行确认。"
    )
    materials_table = document.add_table(rows=1, cols=6)
    materials_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    materials_table.autofit = False
    for index, width in enumerate((0.45, 0.55, 1.35, 1.55, 1.7, 0.55)):
        materials_table.columns[index].width = Inches(width)
    for index, label in enumerate(("编号", "风险", "资料名称", "取证用途", "索取范围", "优先级")):
        _set_cell(materials_table.cell(0, index), label, bold=True, fill=_NAVY, white=True, center=True)
    _repeat_header(materials_table.rows[0])
    for row_index, item in enumerate(draft["materials"], start=1):
        cells = materials_table.add_row().cells
        values = (
            item["id"], item["risk_number"], item["title"], item["purpose"],
            item["requested_scope"], item["priority"],
        )
        for index, value in enumerate(values):
            _set_cell(
                cells[index],
                str(value),
                fill=_PALE_BLUE if row_index % 2 == 0 else "FFFFFF",
                center=index in (0, 1, 5),
            )
    _format_table(materials_table, font_size=8.5)

    document.add_page_break()
    document.add_heading(draft["interview_title"], level=1)
    document.add_paragraph(
        "以下问题用于访谈准备。访谈记录、补充证据和后续判断应另行留痕，不以问题表述替代审计结论。"
    )
    for item in draft["interviews"]:
        heading = document.add_heading(
            f"{item['id']} 关联风险 {item['risk_number']}", level=2
        )
        heading.paragraph_format.keep_with_next = True
        audience = document.add_paragraph()
        audience.paragraph_format.keep_with_next = True
        audience.add_run("建议访谈对象  ").bold = True
        audience.add_run(item["audience"])
        objective = document.add_paragraph()
        objective.paragraph_format.keep_with_next = True
        objective.add_run("访谈目的  ").bold = True
        objective.add_run(item["objective"])
        question = document.add_paragraph()
        question.paragraph_format.space_after = Pt(12)
        question.add_run("访谈问题  ").bold = True
        question.add_run(item["question"])

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer.add_run(
            f"最终草稿 v{draft['version']}  ·  快照 {snapshot['content_sha256'][:12]}"
        )
        footer_run.font.size = Pt(8)
        footer_run.font.color.rgb = _MUTED

    document.save(target)


def _remove_template_marker(document: Document) -> None:
    for paragraph in list(document.paragraphs):
        if paragraph.text.strip() == "[[DOCUMENT_BODY]]":
            element = paragraph._element
            element.getparent().remove(element)


def _set_cell(
    cell,
    text: str,
    *,
    bold: bool = False,
    fill: str = "FFFFFF",
    white: bool = False,
    center: bool = False,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    if white:
        run.font.color.rgb = RGBColor(255, 255, 255)
    shading = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        cell._tc.get_or_add_tcPr().append(shading)
    shading.set(qn("w:fill"), fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell, top=90, start=100, bottom=90, end=100)


def _format_table(table, *, font_size: float) -> None:
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            _set_cell_borders(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.15
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)


def _set_cell_borders(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), _BORDER)


def _set_cell_margins(cell, *, top: int, start: int, bottom: int, end: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)
